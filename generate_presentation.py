import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFont

# Define paths
WORKSPACE_DIR = r"c:\Users\adabayo\OneDrive\Документы\karakol-landing"
IMAGES_DIR = os.path.join(WORKSPACE_DIR, "public", "images")
BRAIN_DIR = r"C:\Users\adabayo\.gemini\antigravity\brain\d1bf0d66-f671-41b3-9652-d2f9601a94d8"

# 1. Copy screenshots
screenshot_hero_src = os.path.join(BRAIN_DIR, "media__1780857531873.png")
screenshot_map_src = os.path.join(BRAIN_DIR, "media__1780857552797.png")
screenshot_hero_dst = os.path.join(IMAGES_DIR, "screenshot_hero.png")
screenshot_map_dst = os.path.join(IMAGES_DIR, "screenshot_map.png")

print("Copying screenshots...")
if os.path.exists(screenshot_hero_src):
    shutil.copy(screenshot_hero_src, screenshot_hero_dst)
if os.path.exists(screenshot_map_src):
    shutil.copy(screenshot_map_src, screenshot_map_dst)

# 2. Draw Competitor Analysis Table
def draw_competitors_table(output_path):
    print("Drawing competitors table image...")
    # 1000x520 canvas with dark navy background (#132a45 - card bg)
    img = Image.new("RGBA", (1000, 520), (19, 42, 69, 255))
    draw = ImageDraw.Draw(img)
    
    try:
        font_header = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 20)
        font_body = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 16)
        font_bold = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 16)
    except Exception as e:
        print("System fonts not loaded, using default:", e)
        font_header = font_body = font_bold = ImageFont.load_default()
        
    # Headers
    headers = [
        "Ресурс", 
        "Интерактивная\nкарта", 
        "Фильтрация\nмаршрутов", 
        "Для соло-\nтуристов", 
        "Дизайн и\nадаптивность"
    ]
    # Column positions
    cols = [30, 260, 440, 620, 800]
    col_widths = [220, 170, 170, 170, 170]
    
    # Draw header row background (#0f2137)
    draw.rectangle([20, 20, 980, 95], fill=(15, 33, 55, 255), outline=(56, 189, 248, 60), width=1)
    
    for i, h in enumerate(headers):
        draw.text((cols[i] + 10, 30), h, font=font_header, fill=(56, 189, 248, 255))
        
    rows = [
        ("Kettik.kg", "no", "no", "no", "warning"),
        ("Visit Kara-Kol", "no", "no", "yes", "no"),
        ("Kyrgyzstan Travel", "no", "no", "no", "warning"),
        ("Karakol Guide (Проект)", "yes", "yes", "yes", "yes")
    ]
    
    y = 105
    for row_idx, r in enumerate(rows):
        # Alternate row backgrounds
        bg_color = (25, 50, 80, 255) if row_idx % 2 == 0 else (19, 42, 69, 255)
        # Highlight our project row
        if row_idx == 3:
            bg_color = (20, 60, 100, 255)
            border_color = (251, 191, 36, 180)
        else:
            border_color = (125, 211, 252, 20)
            
        draw.rectangle([20, y, 980, y + 85], fill=bg_color, outline=border_color, width=2 if row_idx == 3 else 1)
        
        # Resource Name
        res_name = r[0]
        text_color = (251, 191, 36, 255) if row_idx == 3 else (241, 245, 249, 255)
        name_font = font_bold if row_idx == 3 else font_body
        draw.text((cols[0] + 10, y + 32), res_name, font=name_font, fill=text_color)
        
        # Columns 1 to 4 (Checkmarks, Crosses, Warnings)
        for col_idx, status in enumerate(r[1:]):
            cx = cols[col_idx + 1] + col_widths[col_idx + 1] // 2
            cy = y + 42
            if status == "yes":
                # Green Circle
                draw.ellipse([cx - 16, cy - 16, cx + 16, cy + 16], fill=(16, 185, 129, 255))
                # White checkmark
                draw.line([cx - 8, cy, cx - 2, cy + 6], fill=(255, 255, 255, 255), width=3)
                draw.line([cx - 2, cy + 6, cx + 8, cy - 5], fill=(255, 255, 255, 255), width=3)
            elif status == "no":
                # Red Circle
                draw.ellipse([cx - 16, cy - 16, cx + 16, cy + 16], fill=(239, 68, 68, 255))
                # White cross
                draw.line([cx - 6, cy - 6, cx + 6, cy + 6], fill=(255, 255, 255, 255), width=3)
                draw.line([cx + 6, cy - 6, cx - 6, cy + 6], fill=(255, 255, 255, 255), width=3)
            elif status == "warning":
                # Yellow Circle
                draw.ellipse([cx - 16, cy - 16, cx + 16, cy + 16], fill=(245, 158, 11, 255))
                # White exclamation mark
                draw.rectangle([cx - 2, cy - 8, cx + 2, cy + 1], fill=(255, 255, 255, 255))
                draw.ellipse([cx - 2, cy + 4, cx + 2, cy + 8], fill=(255, 255, 255, 255))
                
        y += 95
        
    img.save(output_path)
    print("Table image saved to:", output_path)

# 3. Draw Subjects and Objects Diagram
def draw_subjects_diagram(output_path):
    print("Drawing subjects diagram image...")
    # 1000x550 canvas with dark navy background (#132a45 - card bg)
    img = Image.new("RGBA", (1000, 550), (19, 42, 69, 255))
    draw = ImageDraw.Draw(img)
    
    try:
        font_title = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 20)
        font_body = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 13)
        font_bold = ImageFont.truetype("C:\\Windows\\Fonts\\arial.ttf", 14)
    except Exception as e:
        print("System fonts not loaded, using default:", e)
        font_title = font_body = font_bold = ImageFont.load_default()
        
    # Center node: Karakol Guide
    ccx, ccy = 500, 275
    c_width, c_height = 200, 100
    
    nodes = [
        {"id": "tourists", "title": "Пользователи (Туристы)", "pos": (150, 60), "size": (250, 110), 
         "desc": ["• Самостоятельные туристы", "• Быстрый поиск маршрутов", "• Изучение карты объектов", "• Фильтрация по сложности"]},
         
        {"id": "admin", "title": "Администратор сайта", "pos": (600, 60), "size": (250, 110), 
         "desc": ["• Модерация каталога и контента", "• Добавление достопримечательностей", "• Проверка актуальности информации", "• Обработка форм обратной связи"]},
         
        {"id": "partners", "title": "Местные гиды и партнёры", "pos": (150, 380), "size": (250, 110), 
         "desc": ["• Размещение местных активностей", "• Предоставление услуг гидов", "• Логистика (транспорт, ночлег)", "• Партнёрство с проектом"]},
         
        {"id": "developer", "title": "Разработчик платформы", "pos": (600, 380), "size": (250, 110), 
         "desc": ["• Развертывание и хостинг", "• Техническое обслуживание", "• Разработка нового функционала", "• Оптимизация баз данных (Supabase)"]}
    ]
    
    # Draw connections (Center to Nodes)
    for n in nodes:
        nx, ny = n["pos"]
        nw, nh = n["size"]
        ncx, ncy = nx + nw // 2, ny + nh // 2
        
        # Draw connection line
        draw.line([ncx, ncy, ccx, ccy], fill=(56, 189, 248, 60), width=3)
        draw.ellipse([ncx - 5, ncy - 5, ncx + 5, ncy + 5], fill=(56, 189, 248, 255))
        
    # Draw Center Node
    draw.rectangle([ccx - c_width//2, ccy - c_height//2, ccx + c_width//2, ccy + c_height//2], 
                   fill=(10, 22, 40, 255), outline=(251, 191, 36, 255), width=3)
    draw.text((ccx - 68, ccy - 30), "Karakol Guide", font=font_title, fill=(251, 191, 36, 255))
    draw.text((ccx - 85, ccy), "Объект исследования:\nСистема организации\nтур-информации", font=font_body, fill=(241, 245, 249, 255), align="center")
    
    # Draw Sub-nodes
    for n in nodes:
        nx, ny = n["pos"]
        nw, nh = n["size"]
        # Outer box
        draw.rectangle([nx, ny, nx + nw, ny + nh], fill=(15, 33, 55, 255), outline=(56, 189, 248, 120), width=1)
        # Title bar
        draw.rectangle([nx, ny, nx + nw, ny + 30], fill=(20, 50, 90, 255))
        # Title text
        draw.text((nx + 10, ny + 7), n["title"], font=font_bold, fill=(56, 189, 248, 255))
        # Bullet lines
        by = ny + 36
        for line in n["desc"]:
            draw.text((nx + 10, by), line, font=font_body, fill=(241, 245, 249, 255))
            by += 17
            
    img.save(output_path)
    print("Subjects diagram saved to:", output_path)

# Run image generations
competitors_img_path = os.path.join(IMAGES_DIR, "competitors_analysis.png")
subjects_img_path = os.path.join(IMAGES_DIR, "subjects_diagram.png")
draw_competitors_table(competitors_img_path)
draw_subjects_diagram(subjects_img_path)


# 4. Generate PPTX presentation
prs = Presentation()
# Set aspect ratio to widescreen (16:9)
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color constants
C_NAVY_BG = RGBColor(10, 22, 40)
C_CARD_BG = RGBColor(19, 42, 69)
C_YELLOW = RGBColor(251, 191, 36)
C_BLUE = RGBColor(56, 189, 248)
C_WHITE = RGBColor(241, 245, 249)
C_GRAY = RGBColor(148, 163, 184)

def add_title(slide, text):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Georgia"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = C_YELLOW
    
    # Yellow underline accent
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(1.5), Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_YELLOW
    bar.line.fill.background()

def add_slide_base(title, slide_num):
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    
    # Solid background rectangle
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY_BG
    bg.line.fill.background()
    
    add_title(slide, title)
    add_footer(slide, slide_num)
    return slide

def add_footer(slide, slide_num):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
    tf = footer_box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"Дипломный проект  |  Разработка гида Karakol Guide  |  Слайд {slide_num}"
    p.font.name = "Arial"
    p.font.size = Pt(10)
    p.font.color.rgb = C_GRAY

def add_bullet_points(slide, left, top, width, height, points):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    for i, pt in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        p.level = 0
        
        if ":" in pt and not pt.startswith("http"):
            parts = pt.split(":", 1)
            run1 = p.add_run()
            run1.text = parts[0] + ":"
            run1.font.bold = True
            run1.font.name = "Arial"
            run1.font.size = Pt(14)
            run1.font.color.rgb = C_YELLOW
            
            run2 = p.add_run()
            run2.text = parts[1]
            run2.font.bold = False
            run2.font.name = "Arial"
            run2.font.size = Pt(14)
            run2.font.color.rgb = C_WHITE
        else:
            run = p.add_run()
            run.text = pt
            run.font.name = "Arial"
            run.font.size = Pt(14)
            run.font.color.rgb = C_WHITE

# ----------------- Slide 1: Title Slide -----------------
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY_BG
bg1.line.fill.background()

top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
top_bar.fill.solid()
top_bar.fill.fore_color.rgb = C_YELLOW
top_bar.line.fill.background()

title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.0))
tf = title_box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p = tf.paragraphs[0]
p.text = "РАЗРАБОТКА ИНТЕРАКТИВНОГО ТУРИСТИЧЕСКОГО ГИДА"
p.font.name = "Georgia"
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = C_WHITE

p2 = tf.add_paragraph()
p2.text = "«KARAKOL GUIDE»"
p2.font.name = "Georgia"
p2.font.size = Pt(44)
p2.font.bold = True
p2.font.color.rgb = C_YELLOW
p2.space_before = Pt(10)

sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.8))
tf_sub = sub_box.text_frame
tf_sub.word_wrap = True
tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
p3 = tf_sub.paragraphs[0]
p3.text = "Выпускная квалификационная работа (Дипломный проект)"
p3.font.name = "Arial"
p3.font.size = Pt(18)
p3.font.italic = True
p3.font.color.rgb = C_BLUE

div1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.8), Inches(4.0), Inches(0.02))
div1.fill.solid()
div1.fill.fore_color.rgb = C_GRAY
div1.line.fill.background()

info_box = slide1.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(6.0), Inches(1.5))
tf_info = info_box.text_frame
tf_info.word_wrap = True
tf_info.margin_left = tf_info.margin_top = tf_info.margin_right = tf_info.margin_bottom = 0

p_auth = tf_info.paragraphs[0]
p_auth.text = "Выполнил:"
p_auth.font.name = "Arial"
p_auth.font.size = Pt(12)
p_auth.font.color.rgb = C_GRAY
p_auth.space_after = Pt(2)

p_auth_val = tf_info.add_paragraph()
p_auth_val.text = "Студент группы [Укажите группу]\n[ФИО Студента]"
p_auth_val.font.name = "Arial"
p_auth_val.font.size = Pt(14)
p_auth_val.font.bold = True
p_auth_val.font.color.rgb = C_WHITE
p_auth_val.space_after = Pt(15)

p_sup = tf_info.add_paragraph()
p_sup.text = "Научный руководитель:"
p_sup.font.name = "Arial"
p_sup.font.size = Pt(12)
p_sup.font.color.rgb = C_GRAY
p_sup.space_after = Pt(2)

p_sup_val = tf_info.add_paragraph()
p_sup_val.text = "[Ученая степень, ФИО Руководителя]"
p_sup_val.font.name = "Arial"
p_sup_val.font.size = Pt(14)
p_sup_val.font.bold = True
p_sup_val.font.color.rgb = C_WHITE

year_box = slide1.shapes.add_textbox(Inches(11.0), Inches(6.4), Inches(1.5), Inches(0.5))
tf_year = year_box.text_frame
p_yr = tf_year.paragraphs[0]
p_yr.text = "2026"
p_yr.font.name = "Georgia"
p_yr.font.size = Pt(18)
p_yr.font.bold = True
p_yr.font.color.rgb = C_GRAY
p_yr.alignment = PP_ALIGN.RIGHT

# ----------------- Slide 2: Актуальность темы -----------------
slide2 = add_slide_base("1.1 Актуальность темы", 2)
points2 = [
    "Развитие цифровой экономики: Интернет стал главным инструментом для коммуникации, обучения и планирования путешествий.",
    "Туристический потенциал Кыргызстана: Город Каракол — ключевая дестинация Иссык-Кульской области с живописными горами, озером Ала-Куль и памятниками культуры.",
    "Разнообразие видов отдыха: Рост спроса на треккинг, конные туры, гастрономические и исторические экскурсии требует единой платформы.",
    "Проблема планирования: Отсутствие локальных качественных веб-сервисов с интерактивными картами делает самостоятельное планирование сложным.",
    "Цель Karakol Guide: Создание современного веб-сервиса, объединяющего маршруты, фильтры по сложности и детальную интерактивную карту."
]
add_bullet_points(slide2, Inches(0.8), Inches(1.6), Inches(6.0), Inches(4.8), points2)

# Insert mountains.jpg if exists
img_mountains_path = os.path.join(IMAGES_DIR, "mountains.jpg")
if os.path.exists(img_mountains_path):
    # Border card shape
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.6), Inches(5.3), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = C_BLUE
    card.line.width = Pt(1)
    slide2.shapes.add_picture(img_mountains_path, Inches(7.35), Inches(1.75), Inches(5.0), Inches(4.2))

# ----------------- Slide 3: Анализ аналогов -----------------
slide3 = add_slide_base("1.2 Сравнительный анализ аналогов", 3)
points3 = [
    "Kettik.kg: Крупный портал о туризме в КР. Достоинство: много отзывов и контента. Недостаток: сложный интерфейс, разрозненные данные, нет фильтров маршрутов.",
    "Visit Kara-Kol: Локальный сайт. Достоинство: простая структура. Недостаток: ограниченный объем информации, отсутствие интерактивной карты.",
    "Kyrgyzstan Travel: Сайт туроператора. Достоинство: профессиональные фото. Недостаток: ориентирован только на организованные группы; высокая цена."
]
add_bullet_points(slide3, Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8), points3)

if os.path.exists(competitors_img_path):
    slide3.shapes.add_picture(competitors_img_path, Inches(6.9), Inches(1.6), Inches(5.6), Inches(4.3))

# ----------------- Slide 4: Объект и субъекты исследования -----------------
slide4 = add_slide_base("1.3 Объект и субъекты исследования", 4)
points4 = [
    "Объект исследования: Процесс организации туристической информации о городе Каракол и окрестностях с использованием веб-технологий.",
    "Пользователи (Туристы): Ключевые субъекты. Нуждаются в быстром поиске маршрутов, фильтрации и работе с интерактивной картой.",
    "Администратор сайта: Лицо, отвечающее за модерацию контента, обновление маршрутов и проверку актуальности данных.",
    "Дополнительные субъекты: Разработчик (поддержка) и Местные гиды/партнеры (экспертиза, предложение услуг)."
]
add_bullet_points(slide4, Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8), points4)

if os.path.exists(subjects_img_path):
    slide4.shapes.add_picture(subjects_img_path, Inches(6.8), Inches(1.5), Inches(5.8), Inches(4.8))

# ----------------- Slide 5: Техническое задание (ТЗ) — Архитектура -----------------
slide5 = add_slide_base("1.4 Техническое задание: Требования", 5)
points5 = [
    "Каталог маршрутов: Распределение по активностям (треккинг, конные туры, культура, кулинария) с сортировкой по сложности и длительности.",
    "Природный дизайн: Современный UI на базе природных оттенков Кыргызстана (синий — Иссык-Куль, зеленый — луга, белый — ледники).",
    "Адаптивность интерфейса: Корректное отображение на компьютерах, планшетах и смартфонах (Mobile-first подход).",
    "Навигация: Удобная поисковая строка, система категорий, хлебные крошки и быстрая фильтрация."
]
add_bullet_points(slide5, Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8), points5)

if os.path.exists(screenshot_hero_dst):
    card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = C_BLUE
    card.line.width = Pt(1)
    slide5.shapes.add_picture(screenshot_hero_dst, Inches(7.05), Inches(1.75), Inches(5.3), Inches(4.2))

# ----------------- Slide 6: Интерактивная карта -----------------
slide6 = add_slide_base("Интерактивная карта достопримечательностей", 6)
points6 = [
    "Ключевые точки: Отображение озера Ала-Куль, Жети-Огуз, Алтын-Арашан, Дунганской мечети, Троицкого собора и др.",
    "Интерактивность: Просмотр краткой информации при клике на маркер, прокладка нити маршрута.",
    "Инструменты (Leaflet API): Плавное масштабирование и перемещение карты для детального изучения.",
    "Оптимизация: Динамическая подгрузка маркеров и плавная подгрузка тайловых карт OpenStreetMap."
]
add_bullet_points(slide6, Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8), points6)

if os.path.exists(screenshot_map_dst):
    card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = C_BLUE
    card.line.width = Pt(1)
    slide6.shapes.add_picture(screenshot_map_dst, Inches(7.05), Inches(1.75), Inches(5.3), Inches(4.2))

# ----------------- Slide 7: Технологический стек -----------------
# Let's use the custom grid panel rendering we designed
def add_tech_stack_slide(prs):
    slide = add_slide_base("Технологический стек разработки", 7)
    
    col_width = Inches(2.6)
    col_gap = Inches(0.4)
    top = Inches(1.6)
    height = Inches(4.8)
    
    stack = [
        {
            "cat": "Фронтенд ядро",
            "techs": ["HTML5 / CSS3", "JavaScript (ES6+)", "React.js Component Architecture", "React Router SPA"]
        },
        {
            "cat": "Стилизация и UI",
            "techs": ["Tailwind CSS v4", "CSS Variables & Themes", "Playfair Display / Plus Jakarta Sans fonts", "Fluid Grid Layouts"]
        },
        {
            "cat": "Интерактивная карта",
            "techs": ["Leaflet.js Library", "React-Leaflet wrapper", "OpenStreetMap tiles", "Custom Marker Popups", "Geopositioning"]
        },
        {
            "cat": "Бэкенд и сервисы",
            "techs": ["Supabase Cloud DB", "PostgreSQL database", "REST API Integration", "Admin Dashboard panel", "Contact Form Handlers"]
        }
    ]
    
    for idx, s in enumerate(stack):
        left = Inches(0.8) + idx * (col_width + col_gap)
        
        # Panel
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_width, height)
        panel.fill.solid()
        panel.fill.fore_color.rgb = C_CARD_BG
        panel.line.color.rgb = C_BLUE
        panel.line.width = Pt(1)
        
        p_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.2), col_width - Inches(0.3), height - Inches(0.4))
        tf = p_box.text_frame
        tf.word_wrap = True
        
        # Category
        p_cat = tf.paragraphs[0]
        p_cat.text = s["cat"]
        p_cat.font.name = "Georgia"
        p_cat.font.size = Pt(18)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_YELLOW
        p_cat.space_after = Pt(15)
        
        # Tech items
        for t in s["techs"]:
            p_tech = tf.add_paragraph()
            p_tech.text = "• " + t
            p_tech.font.name = "Arial"
            p_tech.font.size = Pt(12)
            p_tech.font.color.rgb = C_WHITE
            p_tech.space_after = Pt(8)

add_tech_stack_slide(prs)

# ----------------- Slide 8: Сложности при разработке -----------------
def add_challenges_slide(prs):
    slide = add_slide_base("Сложности при разработке и решения", 8)
    
    col_width = Inches(3.6)
    col_gap = Inches(0.4)
    top = Inches(1.6)
    height = Inches(4.8)
    
    challenges = [
        {
            "num": "01",
            "title": "Интеграция Leaflet в React SPA",
            "issue": "Конфликт жизненных циклов Leaflet (работа с реальным DOM) и React (Virtual DOM), вызывающий утечки памяти и сбои маркеров.",
            "sol": "Инициализация карты через useRef и useEffect с вызовом map.remove() при размонтировании. Мемоизация координат и маркеров."
        },
        {
            "num": "02",
            "title": "Оптимизация веса изображений",
            "issue": "Качественные фотографии весили до 3-5 МБ, замедляя первую загрузку (LCP > 4.5с) на мобильных устройствах.",
            "sol": "Конвертация в WebP с сжатием без видимых потерь, lazy loading картинок вне экрана и использование адаптивных размеров srcset."
        },
        {
            "num": "03",
            "title": "Адаптивные фильтры и тач-интерфейс",
            "issue": "Карта перекрывала жесты прокрутки страницы на смартфонах; фильтры были перегружены элементами управления.",
            "sol": "Разработка всплывающей (drawer) панели фильтров, оптимизированной под тач-жесты, и блокировка прокрутки карты при открытом меню."
        }
    ]
    
    for idx, c in enumerate(challenges):
        left = Inches(0.8) + idx * (col_width + col_gap)
        
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_width, height)
        panel.fill.solid()
        panel.fill.fore_color.rgb = C_CARD_BG
        panel.line.color.rgb = C_BLUE
        panel.line.width = Pt(1)
        
        p_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), col_width - Inches(0.4), height - Inches(0.4))
        tf = p_box.text_frame
        tf.word_wrap = True
        
        # Number
        p_num = tf.paragraphs[0]
        p_num.text = c["num"]
        p_num.font.name = "Georgia"
        p_num.font.size = Pt(28)
        p_num.font.bold = True
        p_num.font.color.rgb = C_YELLOW
        p_num.space_after = Pt(8)
        
        # Title
        p_title = tf.add_paragraph()
        p_title.text = c["title"]
        p_title.font.name = "Arial"
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = C_BLUE
        p_title.space_after = Pt(12)
        
        # Problem heading
        p_prob_h = tf.add_paragraph()
        p_prob_h.text = "Проблема:"
        p_prob_h.font.name = "Arial"
        p_prob_h.font.size = Pt(12)
        p_prob_h.font.bold = True
        p_prob_h.font.color.rgb = RGBColor(239, 68, 68) # Red
        p_prob_h.space_after = Pt(2)
        
        # Problem body
        p_prob_b = tf.add_paragraph()
        p_prob_b.text = c["issue"]
        p_prob_b.font.name = "Arial"
        p_prob_b.font.size = Pt(11)
        p_prob_b.font.color.rgb = C_WHITE
        p_prob_b.space_after = Pt(12)
        
        # Solution heading
        p_sol_h = tf.add_paragraph()
        p_sol_h.text = "Решение:"
        p_sol_h.font.name = "Arial"
        p_sol_h.font.size = Pt(12)
        p_sol_h.font.bold = True
        p_sol_h.font.color.rgb = RGBColor(16, 185, 129) # Green
        p_sol_h.space_after = Pt(2)
        
        # Solution body
        p_sol_b = tf.add_paragraph()
        p_sol_b.text = c["sol"]
        p_sol_b.font.name = "Arial"
        p_sol_b.font.size = Pt(11)
        p_sol_b.font.color.rgb = C_WHITE

add_challenges_slide(prs)

# ----------------- Slide 9: Тестирование и верификация -----------------
slide9 = add_slide_base("Тестирование и верификация", 9)
points9 = [
    "Функциональное тестирование: Ручное тестирование всех фильтров, сортировок каталога и корректной загрузки маршрутов.",
    "Интерактивность карты: Валидация кликабельности маркеров, всплывающих окон и подгрузки слоёв Leaflet.",
    "Форма обратной связи: Тестирование отправки сообщений, валидации полей почты и имени.",
    "Нефункциональное тестирование:",
    "  • Адаптивность: Проверка верстки на смартфонах и планшетах (Android/iOS).",
    "  • Кроссбраузерность: Тестирование в Chrome, Safari, Firefox и Edge.",
    "  • Оптимизация скорости: Оценка Google Lighthouse (Performance, SEO > 90)."
]
add_bullet_points(slide9, Inches(0.8), Inches(1.6), Inches(6.0), Inches(4.8), points9)

img_church_path = os.path.join(IMAGES_DIR, "church.jpg")
if os.path.exists(img_church_path):
    card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(1.6), Inches(5.3), Inches(4.5))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD_BG
    card.line.color.rgb = C_BLUE
    card.line.width = Pt(1)
    slide9.shapes.add_picture(img_church_path, Inches(7.35), Inches(1.75), Inches(5.0), Inches(4.2))

# ----------------- Slide 10: Заключение -----------------
slide10 = add_slide_base("Заключение и планы развития", 10)
points10 = [
    "Итоги проекта: Разработан полностью функционирующий веб-ресурс «Karakol Guide» для самостоятельных путешественников.",
    "Практическая значимость: Повышение туристической привлекательности Каракола, предоставление детальной интерактивной базы данных.",
    "Планы развития:",
    "  • Интеграция бронирования: Связь с локальными гостевыми домами и гидами.",
    "  • PWA и офлайн-режим: Добавление поддержки GPS-треков в режиме офлайн.",
    "  • Мобильное приложение: Разработка гида на базе кроссплатформенного React Native.",
    "  • Сообщество: Добавление пользовательских обзоров и системы рейтинга маршрутов."
]
add_bullet_points(slide10, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.2), points10)

# Bottom Thank You text
thanks_box = slide10.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.8))
tf_th = thanks_box.text_frame
p_th = tf_th.paragraphs[0]
p_th.text = "СПАСИБО ЗА ВНИМАНИЕ! ГОТОВ ОТВЕТИТЬ НА ВАШИ ВОПРОСЫ."
p_th.font.name = "Georgia"
p_th.font.size = Pt(20)
p_th.font.bold = True
p_th.font.color.rgb = C_YELLOW
p_th.alignment = PP_ALIGN.CENTER

# Save presentation
pptx_path = os.path.join(WORKSPACE_DIR, "Karakol_Guide_Presentation.pptx")
prs.save(pptx_path)
print("Presentation generated successfully at:", pptx_path)
